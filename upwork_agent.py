"""
upwork_agent.py
===============
Core AI agent logic:
 - Loads freelancer profile & tips
 - Reads / writes the single Word document (cover_letters.docx)
 - Calls Groq API (multiple passes if needed)
 - Returns structured cover letter output
"""

import os
import re
import logging
from datetime import datetime
from typing import Optional

import json
import urllib.request
import urllib.error
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from dotenv import load_dotenv

load_dotenv()

# ─────────────────────────── logging ────────────────────────────
LOG_LEVEL = os.getenv("LOG_LEVEL", "INFO").upper()
logging.basicConfig(
    level=getattr(logging, LOG_LEVEL, logging.INFO),
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger("upwork_agent")

# ─────────────────────────── config ─────────────────────────────
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")
GROQ_MODEL   = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")
WORD_FILE    = os.getenv("WORD_FILE", "cover_letters.docx")
GROQ_URL     = "https://api.groq.com/openai/v1/chat/completions"

# ════════════════════════════════════════════════════════════════
#  FREELANCER PROFILE — edit this section to keep your data fresh
# ════════════════════════════════════════════════════════════════
FREELANCER_PROFILE = """
NAME: San (Full-stack Web Developer & AI Integration Specialist)
EXPERIENCE: 15+ years
LOCATION: Ireland
LINKEDIN: https://www.linkedin.com/in/sanchowdhury/
COMPANY PORTFOLIO: https://euitsols.com/our-work/

--- CORE SKILLS ---
WordPress        | 15+ yrs | custom themes, plugins, multisite, WooCommerce, security
Shopify          | Experienced | API integration, theme customization, app troubleshooting
PHP / Laravel    | 10+ yrs | backend, custom apps, API development, scalable architecture
Python & AI      | Active | document summarizer, data extractor, chatbots, automation
HTML/CSS/JS      | Expert | responsive design, frontend debugging
WooCommerce      | Expert | setup, payment gateways, variable products
AWS / cPanel     | Expert | deployment, SSH, DNS, email config
AI Integrations  | OpenAI, Anthropic Claude, Google Gemini, Meta LLaMA, Hugging Face
AI Systems       | RAG, AI Agents, Vector DBs (Pinecone, Weaviate, FAISS)
Automation       | Zapier, Make (Integromat), n8n, Pabbly, Power Automate, webhooks
Security         | OWASP Top 10, GDPR-aware builds, vulnerability assessment
CI/CD & Cloud    | AWS, DigitalOcean, GCP, Vercel, CI/CD pipelines

--- PAST EMPLOYERS / NOTABLE CLIENTS ---
Microsoft, Oracle, Genesys (enterprise-level background)

--- PORTFOLIO — SHOPIFY ---
https://nectis.com/
https://solodome.com/
https://askmummyanddaddy.com/
https://wrightsofhowth.com/
https://terraorganica.co.uk/

--- PORTFOLIO — WORDPRESS (selected) ---
https://serve.no/
https://safecom.ie/
https://temperrestaurant.com/
http://thekitchendoctor.ie
https://theconstructioncompany.ie/
https://themantl.com/
https://paradisoburger.com/
https://sirjasonwinters.com/
https://usasharp.com/
https://tomkearin.com/
https://www.nyxbikes.com/
https://thedoghouseboutique.com/
http://www.coolingsource.com/
https://jontaffer.com/
http://www.grandmarina.com/
http://www.biotropiclabs.com/
http://gullnummer.no/
http://www.squarejellyfish.com/
http://www.paristexas.ie/
https://www.skgemballasje.no/
https://nlf.ie/
http://aoifelifestyle.com/
https://www.irishmalts.com/
http://mammoth.business/
https://www.irishmalts.com/ (WooCommerce)
https://www.nyxbikes.com/
https://www.biotropiclabs.com/

--- PORTFOLIO — FOOD & BEVERAGE ---
https://temperrestaurant.com/
https://noyalondon.com/
https://themantl.com/
https://paradisoburger.com/
https://bigfernanduk.com/
https://grlondon.co.uk/
https://paristexas.ie/
https://tulsi-galway.com/
https://www.irishmalts.com/

--- PORTFOLIO — HEALTHCARE / DENTAL / MEDICAL ---
https://dhakapharmacy.com.bd/
https://tobehanson.com/
https://www.andresdental.com/
https://sunnyvaledentalspecialty.com/
https://www.bhtbergen.no/
https://www.wellmed.no/
https://www.martintaylordentistry.com/
https://www.wildsmileslivermoreorthodontist.com/

--- PORTFOLIO — CORPORATE / PROFESSIONAL ---
https://www.equinox.com/
https://amrop.ie/
https://rightcheck.io/
https://www.openskydata.com/
https://proptechvc.com/
https://krfthouse.com/
https://www.paxiom.com/
https://www.shelterworks.com/
https://nevadalegalgroup.com/
https://floridatagandtitle.com/
https://kemmlit.ie/

--- PORTFOLIO — REAL ESTATE / CONSTRUCTION ---
https://airtightconstruction.com/
https://proptechvc.com/
https://lakeshore-associates.com/
https://www.casalimones.com/
https://crosscountrymortgage.com/
https://www.aandjfencing.com/
https://dtscontractors.com/
https://msi-engineers.com/
https://unitedhomeexperts.com/

--- PORTFOLIO — CLEANING / SERVICES ---
https://solarprotect.oncestaging.com/
https://thecoolingco.com/
https://www.onstarpestcontrol.com/
http://cleaningforce.com.au/

--- PORTFOLIO — MEDIA / FILM ---
https://arrowfilms.ie/
https://assembly.ie/
https://www.commercialproducersireland.com/

--- PORTFOLIO — PHP / LARAVEL ---
https://dhakapharmacy.com.bd/
https://www.icsb.edu.bd/

--- PORTFOLIO — ENTERTAINMENT / LIFESTYLE ---
https://escapeabilitylv.com/
https://sweetours.com/
https://aoifelifestyle.com/

--- PYTHON / AI PROJECTS ---
Document summarizer (Python + AI)
Real estate & hotel data extractor (Python + AI)
Social media data scraper (Python + AI)
Website chatbot (Python + AI)
Automatic product description & SEO generator (Python Flask API + AI)

--- KEY DIFFERENTIATORS ---
- Based in Ireland (EU timezone — great for UK/EU/US overlap)
- Not just build & leave — handles websites long-term so clients can focus on business
- Enterprise background (Microsoft, Oracle, Genesys) applied to SMB projects
- AI-native: builds with real LLMs, not just ChatGPT wrappers
- 4 months FREE post-launch support on projects
- SDLC-driven process (requirements → design → build → QA → deploy → support)
- OWASP & GDPR-aware by default on all builds
"""

# ════════════════════════════════════════════════════════════════
#  PROPOSAL WRITING TIPS — baked into every generation
# ════════════════════════════════════════════════════════════════
PROPOSAL_TIPS = """
--- UPWORK COVER LETTER RULES (follow every single one) ---

1. LENGTH: 150–200 words MAX. Clients read on phones. First screen must stand alone.
2. HOOK FIRST: Open with 1 sentence that mirrors the client's exact problem or goal.
   Bad: "I am a skilled developer with 15 years experience."
   Good: "I can see you need a fast-loading Shopify store that actually converts — I've built this before."
3. PROOF EARLY: Back the hook with one specific result or relevant project.
4. MIRROR LANGUAGE: Use the client's own words from their job post. Shows you read it carefully.
5. HUMAN TONE: Write like you're talking to someone. No buzzwords. No "I am passionate."
   Say: "I'll own your website so you can focus on running your business."
   Not: "I leverage synergistic solutions to optimize deliverables."
6. STRUCTURED MIDDLE: 2–3 bullet points max. Short. Scannable.
7. DONE = ...: Briefly state what "done" looks like in their language.
8. ONE PORTFOLIO LINK: Pick the single most relevant project — not a dump of 10 links.
9. CLOSE WITH A QUESTION: End with one easy question to start a conversation.
   Not: "Looking forward to hearing from you." (passive, generic)
   Yes: "Do you already have a design ready, or do you need that from scratch too?"
10. NEVER start with "I" — start with the client's problem or a direct statement of value.
11. NO AI BUZZWORDS: No "cutting-edge", "leverage", "passionate", "dynamic", "robust solution".
12. PERSONALIZE: Every letter should feel written for this one client, not a template.
13. KEYWORD INTEGRATION (natural, not listed): WordPress, Shopify, Laravel, WooCommerce,
    AI integration, automation, PHP, full-stack, custom development, API integration.
    Use only what's relevant to this specific job — 8–12 keywords max, never listed as a dump.
"""

# ════════════════════════════════════════════════════════════════
#  GROQ API CALLS
# ════════════════════════════════════════════════════════════════

def _groq_call(messages: list[dict], temperature: float = 0.7, max_tokens: int = 1500) -> str:
    """Single Groq API call. Returns text content or raises."""
    if not GROQ_API_KEY:
        raise ValueError("GROQ_API_KEY is not set. Add it to your .env file.")

    logger.info(f"Calling Groq API | model={GROQ_MODEL} | messages={len(messages)}")

    headers = {
        "Authorization": f"Bearer {GROQ_API_KEY}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": GROQ_MODEL,
        "messages": messages,
        "temperature": temperature,
        "max_tokens": max_tokens,
    }

    try:
        body = json.dumps(payload).encode("utf-8")
        req = urllib.request.Request(
            GROQ_URL,
            data=body,
            headers={**headers, "Content-Length": str(len(body))},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=90) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        content = data["choices"][0]["message"]["content"].strip()
        logger.info(f"Groq API response received | chars={len(content)}")
        return content
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8", errors="replace")
        logger.error(f"Groq API HTTP error: {e.code} — {body}")
        raise RuntimeError(f"Groq API error {e.code}: {body}") from e
    except Exception as e:
        logger.error(f"Groq API call failed: {e}")
        raise


def analyse_job(job_description: str) -> dict:
    """
    PASS 1 — Analyse the job post.
    Extracts: keyword, industry, tech stack, client pain point, tone.
    Returns a dict with structured analysis.
    """
    logger.info("Pass 1: Analysing job description...")

    system = """You are an expert Upwork proposal strategist.
Analyse the job description and return ONLY a JSON object with these keys:
- "keyword": 2-4 word phrase describing the job type (e.g. "WordPress Website Troubleshooting")
- "industry": industry or niche (e.g. "Healthcare", "E-commerce", "SaaS", "Real Estate")
- "tech_stack": list of technologies mentioned or implied (e.g. ["WordPress", "WooCommerce"])
- "client_pain_point": one sentence — what is the client's core problem or goal?
- "client_tone": "formal" or "casual"
- "urgency": "high" or "normal"
- "relevant_portfolio_category": one of these: shopify | wordpress | healthcare | food | corporate | real_estate | services | laravel | ai | media | entertainment
Return ONLY valid JSON. No markdown. No extra text."""

    user = f"Job Description:\n\n{job_description}"

    raw = _groq_call([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ], temperature=0.2, max_tokens=400)

    # Safe JSON parse
    try:
        # Strip possible markdown fences
        clean = re.sub(r"```json|```", "", raw).strip()
        result = json.loads(clean)
        logger.info(f"Job analysis: keyword='{result.get('keyword')}' industry='{result.get('industry')}'")
        return result
    except Exception as e:
        logger.error(f"Failed to parse job analysis JSON: {e} | raw={raw}")
        return {
            "keyword": "General Web Development",
            "industry": "General",
            "tech_stack": [],
            "client_pain_point": job_description[:200],
            "client_tone": "casual",
            "urgency": "normal",
            "relevant_portfolio_category": "wordpress",
        }


def pick_portfolio_links(analysis: dict) -> str:
    """
    PASS 2 — Pick the most relevant 1–3 portfolio links based on job analysis.
    Returns a short string of links with context.
    """
    logger.info("Pass 2: Selecting relevant portfolio links...")

    system = f"""You are a portfolio curator for a freelance web developer.
Based on the job analysis below and the full portfolio, pick the 1 to 3 MOST relevant project URLs.
Return ONLY a short plain-text list like:
• [site description] — https://url.com
• [site description] — https://url.com

Pick sites that match the client's industry and tech stack. 
If the job involves AI, pick an AI project or SaaS site.
If it's a restaurant, pick food/beverage sites.
If it's healthcare, pick medical sites.
Always prefer specificity over breadth.

FULL PORTFOLIO:
{FREELANCER_PROFILE}"""

    user = f"""Job Analysis:
Industry: {analysis.get('industry')}
Tech Stack: {analysis.get('tech_stack')}
Category hint: {analysis.get('relevant_portfolio_category')}
Pain point: {analysis.get('client_pain_point')}"""

    result = _groq_call([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ], temperature=0.3, max_tokens=200)

    logger.info(f"Portfolio selected:\n{result}")
    return result


def generate_cover_letter(
    job_description: str,
    analysis: dict,
    portfolio_links: str,
    past_letters: Optional[str] = None,
) -> str:
    """
    PASS 3 — Generate the actual cover letter.
    Uses all context: profile, tips, analysis, portfolio, and past letters for this keyword.
    """
    logger.info("Pass 3: Generating cover letter...")

    past_context = ""
    if past_letters:
        past_context = f"""
--- PAST COVER LETTERS FOR THIS KEYWORD (use as reference — do NOT copy, but learn the pattern) ---
{past_letters}
---"""

    system = f"""You are an expert Upwork proposal writer who writes proposals that win jobs.
You write for a specific freelancer. Here is everything about them:

{FREELANCER_PROFILE}

{PROPOSAL_TIPS}

Your job: Write a winning Upwork cover letter for the job description the user provides.
{past_context}

STRICT RULES:
- 150–200 words MAXIMUM
- Never start with "I"
- No buzzwords, no robotic language
- Sound like a real human talking — not a CV
- Use the client's language from their job post
- Pick 1 specific portfolio link (already selected for you — use it)
- End with one easy conversational question
- Integrate relevant keywords naturally (never list them)
- Include the "I'll handle it so you can focus on your business" angle where it fits naturally
- Leave one placeholder in square brackets if client name is unknown: [Client Name] or [Hi there]"""

    user = f"""Write a cover letter for this job:

JOB DESCRIPTION:
{job_description}

JOB ANALYSIS (use this to guide your writing):
- Keyword: {analysis.get('keyword')}
- Industry: {analysis.get('industry')}
- Tech Stack: {', '.join(analysis.get('tech_stack', []))}
- Client's core pain point: {analysis.get('client_pain_point')}
- Tone: {analysis.get('client_tone')}
- Urgency: {analysis.get('urgency')}

MOST RELEVANT PORTFOLIO LINKS TO USE (pick 1 max):
{portfolio_links}

Now write the cover letter. Human, direct, no fluff."""

    letter = _groq_call([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ], temperature=0.75, max_tokens=600)

    logger.info(f"Cover letter generated | chars={len(letter)}")
    return letter


def polish_letter(draft: str, analysis: dict) -> str:
    """
    PASS 4 — Quick polish pass.
    Checks word count, removes AI-isms, ensures it ends with a question.
    """
    logger.info("Pass 4: Polishing cover letter...")

    system = """You are an editor who makes Upwork cover letters sound genuinely human.
Your job: take the draft and make minimal improvements only if needed.

Check for:
1. Does it start with "I"? If yes, rework the opening.
2. Are there any AI buzzwords? (leverage, cutting-edge, robust, passionate, dynamic, synergistic) — remove them.
3. Does it end with a conversational question? If not, add one.
4. Is it over 210 words? If yes, trim — cut fluff, not substance.
5. Does it feel like a human wrote it, or a robot? Fix any robotic phrases.

Return ONLY the final polished cover letter text. No commentary. No labels."""

    user = f"""Draft cover letter:

{draft}

Job industry: {analysis.get('industry')}
Client tone: {analysis.get('client_tone')}

Polish it. Keep it real."""

    polished = _groq_call([
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ], temperature=0.4, max_tokens=600)

    logger.info(f"Polished letter | chars={len(polished)}")
    return polished


# ════════════════════════════════════════════════════════════════
#  WORD DOCUMENT — single source of truth
# ════════════════════════════════════════════════════════════════

def _heading_style(para, level: int = 1):
    """Apply simple heading formatting."""
    run = para.runs[0] if para.runs else para.add_run()
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    elif level == 2:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)


def read_past_letters(keyword: str) -> Optional[str]:
    """
    Read the Word doc and find any saved letters matching this keyword.
    Returns the text of past letters or None.
    """
    if not os.path.exists(WORD_FILE):
        logger.info(f"Word file '{WORD_FILE}' does not exist yet — no past letters.")
        return None

    logger.info(f"Reading past letters for keyword: '{keyword}' from '{WORD_FILE}'")

    try:
        doc = Document(WORD_FILE)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        # Find sections that match the keyword (case-insensitive)
        keyword_lower = keyword.lower()
        sections = []
        current_section = []
        in_section = False

        for para in doc.paragraphs:
            text = para.text.strip()
            if "KEYWORD:" in text and keyword_lower in text.lower():
                in_section = True
                current_section = [text]
            elif "KEYWORD:" in text and in_section:
                # New keyword section started — save previous
                sections.append("\n".join(current_section))
                current_section = [text]
                in_section = False
            elif in_section:
                current_section.append(text)

        if in_section and current_section:
            sections.append("\n".join(current_section))

        if sections:
            logger.info(f"Found {len(sections)} past letter(s) for keyword '{keyword}'")
            return "\n\n---\n\n".join(sections[-2:])  # Return last 2 at most
        else:
            logger.info(f"No past letters found for keyword '{keyword}'")
            return None

    except Exception as e:
        logger.error(f"Error reading Word file: {e}")
        return None


def save_to_word(keyword: str, job_description: str, letter: str, analysis: dict) -> str:
    """
    Append a new cover letter entry to the Word document.
    Creates the file if it doesn't exist.
    Returns the file path.
    """
    logger.info(f"Saving cover letter to '{WORD_FILE}' under keyword '{keyword}'")

    try:
        if os.path.exists(WORD_FILE):
            doc = Document(WORD_FILE)
            logger.info(f"Appending to existing Word file")
        else:
            doc = Document()
            # Document title
            title = doc.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = title.add_run("Upwork Cover Letter Library")
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            doc.add_paragraph()
            logger.info(f"Creating new Word file '{WORD_FILE}'")

        # Page break between entries (except first)
        if len(doc.paragraphs) > 3:
            doc.add_page_break()

        # ── Entry Header ──
        h = doc.add_paragraph()
        run = h.add_run(f"KEYWORD: {keyword}")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

        # Metadata
        meta = doc.add_paragraph()
        meta.add_run(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
                     f"Industry: {analysis.get('industry', 'N/A')}  |  "
                     f"Tech: {', '.join(analysis.get('tech_stack', []))}")
        meta.runs[0].font.size = Pt(9)
        meta.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)

        doc.add_paragraph()

        # ── Job Description ──
        jd_label = doc.add_paragraph()
        run = jd_label.add_run("JOB DESCRIPTION")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

        jd_para = doc.add_paragraph(job_description.strip())
        jd_para.runs[0].font.size = Pt(9)
        jd_para.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

        doc.add_paragraph()

        # ── Cover Letter ──
        cl_label = doc.add_paragraph()
        run = cl_label.add_run("COVER LETTER")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)

        # Split by paragraphs to preserve structure
        for para_text in letter.strip().split("\n"):
            p = doc.add_paragraph(para_text)
            if para_text.strip():
                p.runs[0].font.size = Pt(10.5)

        # Divider
        doc.add_paragraph()
        divider = doc.add_paragraph("─" * 80)
        divider.runs[0].font.size = Pt(8)
        divider.runs[0].font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

        doc.save(WORD_FILE)
        logger.info(f"Saved successfully to '{WORD_FILE}'")
        return os.path.abspath(WORD_FILE)

    except Exception as e:
        logger.error(f"Error saving to Word file: {e}")
        raise


# ════════════════════════════════════════════════════════════════
#  MAIN ORCHESTRATOR
# ════════════════════════════════════════════════════════════════

def generate(job_description: str) -> dict:
    """
    Full pipeline:
      1. Analyse job
      2. Read past letters for this keyword
      3. Pick portfolio links
      4. Generate cover letter
      5. Polish cover letter
      6. Save to Word doc
    Returns dict with: keyword, letter, word_file, analysis
    """
    logger.info("=" * 60)
    logger.info("Starting cover letter generation pipeline")
    logger.info("=" * 60)

    if not job_description.strip():
        raise ValueError("Job description cannot be empty.")

    # Pass 1 — analyse
    analysis = analyse_job(job_description)
    keyword = analysis.get("keyword", "General Web Development")

    # Read past letters for context
    past_letters = read_past_letters(keyword)

    # Pass 2 — pick portfolio
    portfolio_links = pick_portfolio_links(analysis)

    # Pass 3 — generate
    draft = generate_cover_letter(job_description, analysis, portfolio_links, past_letters)

    # Pass 4 — polish
    final_letter = polish_letter(draft, analysis)

    # Save to Word
    word_path = save_to_word(keyword, job_description, final_letter, analysis)

    logger.info("Pipeline complete ✓")
    logger.info("=" * 60)

    return {
        "keyword":      keyword,
        "industry":     analysis.get("industry", ""),
        "tech_stack":   analysis.get("tech_stack", []),
        "letter":       final_letter,
        "portfolio":    portfolio_links,
        "word_file":    word_path,
        "has_past":     past_letters is not None,
    }
